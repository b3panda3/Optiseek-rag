"""docx_parser.py — Extract text from Word documents.

Uses python-docx to extract paragraphs and table cell text. Table text is
included because challenge questions may reference data held in tables
(e.g., the roadmap Q3 FY27 entry is in a table cell).
"""

from __future__ import annotations

import logging
from pathlib import Path

log = logging.getLogger("optiseek.rag.parsers.docx")


def parse_docx(path: Path, corpus_root: Path) -> list[dict] | None:
    """Extract text from a .docx file. Returns one document with all text."""
    from docx import Document

    rel_path = str(path.relative_to(corpus_root))

    try:
        doc = Document(str(path))
    except Exception as e:
        log.warning("DOCX open failed: %s: %s", rel_path, e)
        return None

    parts: list[str] = []

    # Extract paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Extract tables (each row joined by tabs, rows joined by newlines)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))

    full_text = "\n".join(parts)
    if not full_text.strip():
        return None

    return [{
        "file": rel_path,
        "text": full_text,
        "metadata": {
            "type": "docx",
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
        },
    }]
